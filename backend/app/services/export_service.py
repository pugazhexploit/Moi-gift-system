"""Safe in-memory report exports for CSV, Excel, and PDF downloads."""

from __future__ import annotations

import asyncio
import csv
import io
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from html import escape
from typing import Any

from app.core.exceptions import AppError
from app.schemas.report import ExportFormat


@dataclass(frozen=True)
class ExportFile:
    content: bytes
    filename: str
    media_type: str


class ExportService:
    MAX_PDF_ROWS = 1_000

    async def build(
        self, export_format: ExportFormat, title: str, columns: list[str], rows: list[list[Any]]
    ) -> ExportFile:
        if export_format is ExportFormat.CSV:
            return await asyncio.to_thread(self._csv, title, columns, rows)
        if export_format is ExportFormat.XLSX:
            return await asyncio.to_thread(self._xlsx, title, columns, rows)
        if len(rows) > self.MAX_PDF_ROWS:
            raise AppError("EXPORT_LIMIT_EXCEEDED", "PDF exports are limited to 1000 rows", 422)
        return await asyncio.to_thread(self._pdf, title, columns, rows)

    @classmethod
    def _display_value(cls, value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.isoformat()
        return str(value)

    @classmethod
    def _safe_text(cls, value: Any) -> Any:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        return value

    def _csv(self, title: str, columns: list[str], rows: list[list[Any]]) -> ExportFile:
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow(columns)
        for row in rows:
            writer.writerow(
                [
                    self._safe_text(self._display_value(value))
                    if isinstance(value, str)
                    else self._display_value(value)
                    for value in row
                ]
            )
        return ExportFile(
            content=buffer.getvalue().encode("utf-8-sig"),
            filename=f"{title}.csv",
            media_type="text/csv; charset=utf-8",
        )

    def _xlsx(self, title: str, columns: list[str], rows: list[list[Any]]) -> ExportFile:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise AppError("EXPORT_UNAVAILABLE", "Excel export is unavailable", 503) from exc

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Report"
        worksheet.freeze_panes = "A2"
        worksheet.append(columns)
        header_fill = PatternFill("solid", fgColor="0F766E")
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for row in rows:
            worksheet.append([self._safe_text(value) for value in row])
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            max_width = max(len(self._display_value(cell.value)) for cell in column_cells)
            worksheet.column_dimensions[column_letter].width = min(max(max_width + 2, 12), 40)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, Decimal):
                    cell.number_format = "#,##0.00"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
        buffer = io.BytesIO()
        workbook.save(buffer)
        return ExportFile(
            content=buffer.getvalue(),
            filename=f"{title}.xlsx",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def _pdf(self, title: str, columns: list[str], rows: list[list[Any]]) -> ExportFile:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise AppError("EXPORT_UNAVAILABLE", "PDF export is unavailable", 503) from exc

        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
        )
        styles = getSampleStyleSheet()
        body_style = styles["BodyText"]
        body_style.fontSize = 7
        body_style.leading = 9
        table_data = [columns]
        for row in rows:
            table_data.append(
                [Paragraph(escape(self._display_value(value)), body_style) for value in row]
            )
        available_width = landscape(A4)[0] - 24 * mm
        table = Table(table_data, colWidths=[available_width / len(columns)] * len(columns), repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("LEADING", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ]
            )
        )
        document.build([Paragraph(title.replace("_", " ").title(), styles["Title"]), Spacer(1, 8), table])
        return ExportFile(content=buffer.getvalue(), filename=f"{title}.pdf", media_type="application/pdf")
